"""
Pure logic for the weekday-morning 07:00 gateway Check Now.

Three pieces:
  * next_morning_check_time(now)        — next Taipei 07:00 that lands on an
                                          NYSE trading day (strictly future;
                                          weekends + NYSE holidays skipped)
  * format_morning_check_message(names) — Telegram alert body, or None when
                                          nothing is offline (all gateways
                                          healthy → caller stays silent)
  * constituents_staleness_warning(...) — Telegram warning when the ADR-0002
                                          SP500 constituents snapshot is
                                          stale, or None when fresh
                                          (read_constituents_last_date is its
                                          I/O companion — the one deliberate
                                          exception to "this module is pure")

Side effects (real IB re-probe, Tk after, Telegram send) live in
control_panel. This module stays pure so the scheduling and message logic
can be exercised in milliseconds without any infrastructure.

Note on the calendar test: we feed `is_nyse_trading_day` the *Taipei*
calendar date (same convention as ty_check). At 07:00 Taipei that maps to
the previous US session's evening, but the intent here is simply "run a
connection health-check on each Taipei trading weekday, skip weekends and
US holidays" — not to align with a specific NYSE session — so the Taipei
date is the right and simplest key.
"""

from __future__ import annotations

from datetime import date, datetime, time, timedelta
from pathlib import Path

from market_calendar import is_nyse_trading_day

# Trigger time-of-day in Taipei.
_TRIGGER_TIME = time(7, 0)

# Cap forward scan so a corrupt calendar can't infinite-loop us.
_MAX_SCAN_DAYS = 14

# Seconds to wait after firing Check Now before reading is_online state.
# Must exceed gateway_monitor.CONNECT_TIMEOUT_SEC (10) so a re-spawned
# monitor thread's fresh connect attempt has time to either succeed or give
# up before we decide whether to alert. Mirrors rosa_poll.RETRY_WAIT_SEC.
MORNING_SETTLE_SEC = 12.0


def next_morning_check_time(now_taipei: datetime) -> datetime:
    """Return the next Taipei datetime (strictly future) at which we should
    run the morning gateway check. Always 07:00, skipping any day that is a
    weekend or NYSE holiday."""
    candidate = now_taipei.replace(
        hour=_TRIGGER_TIME.hour, minute=_TRIGGER_TIME.minute,
        second=0, microsecond=0,
    )
    if candidate <= now_taipei:
        candidate = candidate + timedelta(days=1)
        candidate = candidate.replace(
            hour=_TRIGGER_TIME.hour, minute=_TRIGGER_TIME.minute,
            second=0, microsecond=0,
        )

    for _ in range(_MAX_SCAN_DAYS):
        if is_nyse_trading_day(candidate.date()):
            return candidate
        candidate = candidate + timedelta(days=1)

    raise RuntimeError(
        f"next_morning_check_time: no NYSE trading day found within "
        f"{_MAX_SCAN_DAYS} days from {now_taipei.isoformat()}"
    )


def format_morning_check_message(offline_names: list[str]) -> str | None:
    """Telegram alert body for the morning check, or None when the offline
    list is empty (all gateways healthy → caller sends nothing).

    Listing the offline accounts (not the healthy ones) keeps the alert
    actionable: it names exactly which Gateways need a manual restart."""
    if not offline_names:
        return None
    lines = "\n".join(
        f"\U0001F534 <b>{name}</b> Gateway offline" for name in offline_names
    )
    return (
        f"⚠️ <b>Morning check 07:00</b>\n{lines}\n"
        f"<code>Check Now re-probe failed — manual Gateway restart needed</code>"
    )


# ── ADR-0002 constituents staleness reminder ────────────────────────
# Refresh stays on-demand (auto-fetch rejected: Wikipedia vandalism / schema
# risk must not silently feed live decisions). This reminder exists because
# on-demand was forgotten for 5 weeks (2026-06/07) and stale membership
# re-creates the CIEN misclassification class of bug. Reminder only — it
# never fetches.

CONSTITUENTS_MAX_AGE_DAYS = 7

_REFRESH_CMD = "python data/refresh_constituents.py"


def read_constituents_last_date(xlsx_path: Path) -> date | None:
    """I/O companion for constituents_staleness_warning: newest dense-snapshot
    date in the ADR-0002 constituents file.

    Sheet layout mirrors SelectionEngine._load_constituents: snapshots live in
    digit-named year sheets ('2015'…'2026', column A = Date); 'Summary' /
    'Counts' are metadata. We scan column A of the newest year sheet.

    Returns None when the file is missing or unreadable — halt-preference
    convention (ADR-0003): callers alert on None instead of staying silent."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(xlsx_path, read_only=True)
        try:
            year_sheets = sorted(s for s in wb.sheetnames if s.isdigit())
            if not year_sheets:
                return None
            ws = wb[year_sheets[-1]]
            last: date | None = None
            for (v,) in ws.iter_rows(min_col=1, max_col=1, values_only=True):
                if isinstance(v, datetime):
                    last = v.date()
                elif isinstance(v, date):
                    last = v
            return last
        finally:
            wb.close()
    except Exception:
        return None


def constituents_staleness_warning(
    last_snapshot: date | None,
    today: date,
    max_age_days: int = CONSTITUENTS_MAX_AGE_DAYS,
) -> str | None:
    """Telegram warning body when the SP500 constituents snapshot is older
    than max_age_days (or unreadable), None when fresh enough.

    Age is measured in calendar days: the refresh cadence ADR-0002 suggests
    is "before the weekly Wednesday SP500 rebalance", so 7 calendar days is
    one missed cycle — warn from the first missed week onward."""
    if last_snapshot is None:
        return (
            "⚠️ <b>SP500 constituents</b> 快照日期無法讀取（檔案缺失或格式異常）"
            f"— 請跑 <code>{_REFRESH_CMD}</code>"
        )
    age_days = (today - last_snapshot).days
    if age_days <= max_age_days:
        return None
    return (
        f"⚠️ <b>SP500 constituents</b> 已 {age_days} 天未更新"
        f"（最後快照 {last_snapshot.isoformat()}）"
        f"— 請跑 <code>{_REFRESH_CMD}</code>"
    )
