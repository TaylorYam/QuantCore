"""Audit and stage verified historical Constituents membership repairs (#54).

The January 2026 workbook was reconstructed from a current ticker list.  That
made later joiners appear under their current symbols before they joined the
index.  This tool applies only the reviewed effective dates in
``constituents_history_repairs.json``:

* ``audit`` reports pre-join rows without changing anything.
* ``repair`` writes a new workbook and leaves the source untouched.

No old ticker is invented, no price file is renamed, and today's constituents
are never used to infer historical membership.  The repaired workbook must be
deployed separately after review; this module never writes the live reference
file in place.
"""
from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
MANIFEST_PATH = Path(__file__).with_name("constituents_history_repairs.json")


@dataclass(frozen=True)
class RepairRule:
    symbol: str
    effective_date: date
    joined_as: str
    evidence: str


def _parse_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    try:
        return date.fromisoformat(text[:10])
    except ValueError as exc:
        raise ValueError(f"invalid ISO date: {value!r}") from exc


def _normalize_symbol(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().upper()
    return text[:-4] if text.endswith(".TXT") else text


def load_rules(path: Path = MANIFEST_PATH) -> tuple[dict, tuple[RepairRule, ...]]:
    """Load and validate the reviewed repair manifest."""
    manifest = json.loads(Path(path).read_text(encoding="utf-8"))
    source = manifest.get("source") or {}
    if not source.get("url") or not source.get("revision_id"):
        raise ValueError("repair manifest source must include url and revision_id")

    rules = []
    seen = set()
    for raw in manifest.get("rules", []):
        symbol = _normalize_symbol(raw.get("symbol"))
        joined_as = _normalize_symbol(raw.get("joined_as"))
        if not symbol or symbol in seen:
            raise ValueError(f"duplicate or empty repair symbol: {symbol!r}")
        if not joined_as or not raw.get("evidence"):
            raise ValueError(f"repair rule for {symbol} is incomplete")
        seen.add(symbol)
        rules.append(RepairRule(symbol, _parse_date(raw.get("effective_date")),
                                joined_as, str(raw["evidence"])))
    if not rules:
        raise ValueError("repair manifest has no rules")
    return manifest, tuple(rules)


def _year_sheets(workbook) -> list:
    return [sheet for sheet in workbook.worksheets if sheet.title.isdigit()]


def _columns(sheet) -> tuple[int, list[int]]:
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    date_columns = [i for i, value in enumerate(headers) if str(value).strip() == "Date"]
    if len(date_columns) != 1:
        raise ValueError(f"{sheet.title}: expected one Date column")
    date_column = date_columns[0]
    return date_column, [i for i in range(len(headers)) if i != date_column]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _base_record(rule: RepairRule, manifest: dict) -> dict:
    return {
        "symbol": rule.symbol,
        "effective_date": rule.effective_date.isoformat(),
        "joined_as": rule.joined_as,
        "evidence": rule.evidence,
        "source": manifest["source"],
        "pre_join_rows": 0,
        "pre_join_cells": 0,
        "first_false_date": None,
        "last_false_date": None,
        "duplicate_dates": [],
        "price_bar_rows": None,
        "price_bar_first_date": None,
        "price_bar_last_date": None,
    }


def _price_dates(data_root: Path, symbol: str) -> set[date]:
    path = Path(data_root) / "stocks" / f"{symbol}.csv"
    if not path.is_file():
        return set()
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames or "Date" not in reader.fieldnames:
            return set()
        dates = set()
        for row in reader:
            try:
                dates.add(_parse_date(row.get("Date", "")))
            except ValueError:
                continue
        return dates


def _audit_open_workbook(workbook, manifest: dict, rules: Iterable[RepairRule],
                         data_root: Path | None = None) -> dict:
    rule_list = tuple(rules)
    by_symbol = {rule.symbol: rule for rule in rule_list}
    records = {rule.symbol: _base_record(rule, manifest) for rule in rule_list}
    dates_by_symbol: dict[str, dict[date, int]] = {rule.symbol: {} for rule in rule_list}
    invalid_date_rows = 0

    for sheet in _year_sheets(workbook):
        date_column, value_columns = _columns(sheet)
        for row in sheet.iter_rows(min_row=2):
            when_value = row[date_column].value
            try:
                when = _parse_date(when_value)
            except ValueError:
                invalid_date_rows += 1
                continue
            for column in value_columns:
                symbol = _normalize_symbol(row[column].value)
                rule = by_symbol.get(symbol)
                if rule is None or when >= rule.effective_date:
                    continue
                counts = dates_by_symbol[symbol]
                counts[when] = counts.get(when, 0) + 1

    for rule in rule_list:
        record = records[rule.symbol]
        counts = dates_by_symbol[rule.symbol]
        false_dates = sorted(counts)
        record["pre_join_rows"] = len(false_dates)
        record["pre_join_cells"] = sum(counts.values())
        record["first_false_date"] = false_dates[0].isoformat() if false_dates else None
        record["last_false_date"] = false_dates[-1].isoformat() if false_dates else None
        record["duplicate_dates"] = [day.isoformat() for day in false_dates if counts[day] > 1]
        if data_root is not None:
            bars = sorted(set(false_dates) & _price_dates(data_root, rule.symbol))
            record["price_bar_rows"] = len(bars)
            record["price_bar_first_date"] = bars[0].isoformat() if bars else None
            record["price_bar_last_date"] = bars[-1].isoformat() if bars else None

    price_bar_rows = [r["price_bar_rows"] for r in records.values()]
    return {
        "schema_version": 1,
        "source": manifest["source"],
        "rules": list(records.values()),
        "total_false_member_rows": sum(r["pre_join_rows"] for r in records.values()),
        "total_false_member_cells": sum(r["pre_join_cells"] for r in records.values()),
        "invalid_date_rows": invalid_date_rows,
        "data_root": str(Path(data_root).resolve()) if data_root is not None else None,
        "total_price_bar_rows": sum(price_bar_rows) if data_root is not None else None,
        "clean": all(r["pre_join_cells"] == 0 for r in records.values()),
    }


def audit_workbook(path: Path, manifest_path: Path = MANIFEST_PATH,
                   data_root: Path | None = None) -> dict:
    """Return a deterministic report of verified pre-join membership."""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)
    manifest, rules = load_rules(manifest_path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        report = _audit_open_workbook(workbook, manifest, rules, data_root)
    finally:
        workbook.close()
    report["workbook"] = str(path.resolve())
    report["workbook_sha256"] = _sha256(path)
    return report


def repair_workbook(source: Path, output: Path,
                    manifest_path: Path = MANIFEST_PATH,
                    data_root: Path | None = None) -> dict:
    """Write a repaired copy, refusing to overwrite either input or output."""
    source, output = Path(source), Path(output)
    if not source.is_file():
        raise FileNotFoundError(source)
    if source.resolve() == output.resolve():
        raise ValueError("output must be different from source")
    if output.exists():
        raise FileExistsError(f"refusing to overwrite staging output: {output}")

    manifest, rules = load_rules(manifest_path)
    before = audit_workbook(source, manifest_path, data_root)
    by_symbol = {rule.symbol: rule for rule in rules}
    removed: dict[str, int] = {rule.symbol: 0 for rule in rules}

    workbook = load_workbook(source, read_only=False, data_only=False)
    try:
        for sheet in _year_sheets(workbook):
            date_column, value_columns = _columns(sheet)
            for row in sheet.iter_rows(min_row=2):
                try:
                    when = _parse_date(row[date_column].value)
                except ValueError:
                    continue
                for column in value_columns:
                    symbol = _normalize_symbol(row[column].value)
                    rule = by_symbol.get(symbol)
                    if rule is not None and when < rule.effective_date:
                        row[column].value = None
                        removed[symbol] += 1
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output)
    finally:
        workbook.close()

    after = audit_workbook(output, manifest_path, data_root)
    if not after["clean"]:
        raise ValueError("staged workbook still contains pre-join members")
    return {
        "schema_version": 1,
        "source": str(source.resolve()),
        "source_sha256": before["workbook_sha256"],
        "output": str(output.resolve()),
        "output_sha256": after["workbook_sha256"],
        "before": before,
        "removed_cells": removed,
        "removed_total": sum(removed.values()),
        "after": after,
        "rollback": "Delete the staged output; the source workbook was never modified.",
    }


def _write_json(path: Path, value: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)

    audit_parser = sub.add_parser("audit", help="report verified pre-join members")
    audit_parser.add_argument("--source", required=True, type=Path)
    audit_parser.add_argument("--data-root", type=Path,
                              help="optional data root for price-bar overlap evidence")
    audit_parser.add_argument("--manifest", type=Path, default=MANIFEST_PATH)

    repair_parser = sub.add_parser("repair", help="write a new repaired workbook")
    repair_parser.add_argument("--source", required=True, type=Path)
    repair_parser.add_argument("--out", required=True, type=Path)
    repair_parser.add_argument("--report", type=Path)
    repair_parser.add_argument("--data-root", type=Path,
                               help="optional data root for price-bar overlap evidence")
    repair_parser.add_argument("--manifest", type=Path, default=MANIFEST_PATH)

    args = parser.parse_args()
    if args.command == "audit":
        report = audit_workbook(args.source, args.manifest, args.data_root)
    else:
        report = repair_workbook(args.source, args.out, args.manifest, args.data_root)
        if args.report:
            _write_json(args.report, report)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
