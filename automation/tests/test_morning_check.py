"""
Unit tests for morning_check — weekday-morning 07:00 gateway Check Now.

Two pieces under test:
  * next_morning_check_time(now)        → next Taipei 07:00 on an NYSE
                                          trading day (weekends + holidays
                                          skipped)
  * format_morning_check_message(names) → Telegram alert string, or None
                                          when nothing is offline

All pure (no IB, no Tk, no real time).
"""

from __future__ import annotations

from datetime import date, datetime
from zoneinfo import ZoneInfo

import pytest

TAIPEI = ZoneInfo("Asia/Taipei")


def _t(y, mo, d, h, mi=0):
    """Convenience: Taipei datetime defaulting to :00 (the trigger minute)."""
    return datetime(y, mo, d, h, mi, tzinfo=TAIPEI)


# ── next_morning_check_time: weekday math ──────────────────────────
# Calendar anchors (2026): Mon May 4, Tue May 5, Fri May 8, Sat May 9,
# Sun May 10, Mon May 11, Tue May 12 — none are NYSE holidays.

def test_next_from_monday_evening_returns_tue_0700():
    from morning_check import next_morning_check_time
    # Mon May 4 18:00 → today's 07:00 already past → Tue May 5 07:00
    assert next_morning_check_time(_t(2026, 5, 4, 18, 0)) == _t(2026, 5, 5, 7, 0)


def test_next_from_tue_before_trigger_returns_same_day():
    """Before 07:00 on a trading Tue, the answer is today 07:00."""
    from morning_check import next_morning_check_time
    assert next_morning_check_time(_t(2026, 5, 5, 6, 0)) == _t(2026, 5, 5, 7, 0)


def test_next_from_one_minute_before_trigger_returns_same_day():
    from morning_check import next_morning_check_time
    assert next_morning_check_time(_t(2026, 5, 5, 6, 59)) == _t(2026, 5, 5, 7, 0)


def test_next_from_trigger_exact_returns_next_day():
    """Strictly future — exactly on 07:00 means the after() callback already
    consumed today's tick, so the next one is tomorrow."""
    from morning_check import next_morning_check_time
    assert next_morning_check_time(_t(2026, 5, 5, 7, 0)) == _t(2026, 5, 6, 7, 0)


def test_next_from_just_after_trigger_returns_next_day():
    from morning_check import next_morning_check_time
    assert next_morning_check_time(_t(2026, 5, 5, 7, 1)) == _t(2026, 5, 6, 7, 0)


def test_next_from_friday_before_trigger_returns_same_friday():
    from morning_check import next_morning_check_time
    assert next_morning_check_time(_t(2026, 5, 8, 6, 0)) == _t(2026, 5, 8, 7, 0)


def test_next_from_friday_after_trigger_skips_weekend_to_monday():
    from morning_check import next_morning_check_time
    # Fri May 8 08:00 → Sat/Sun skipped → Mon May 11 07:00
    assert next_morning_check_time(_t(2026, 5, 8, 8, 0)) == _t(2026, 5, 11, 7, 0)


def test_next_from_saturday_returns_monday():
    from morning_check import next_morning_check_time
    assert next_morning_check_time(_t(2026, 5, 9, 12, 0)) == _t(2026, 5, 11, 7, 0)


def test_next_from_sunday_returns_monday():
    from morning_check import next_morning_check_time
    assert next_morning_check_time(_t(2026, 5, 10, 23, 0)) == _t(2026, 5, 11, 7, 0)


# ── next_morning_check_time: NYSE holiday skipping ─────────────────

def test_next_skips_holiday_tuesday(monkeypatch):
    """Mock Tue May 5 as a NYSE holiday → next check rolls to Wed May 6."""
    import morning_check

    real = morning_check.is_nyse_trading_day

    def mock_is_trading(d):
        if d == date(2026, 5, 5):
            return False
        return real(d)

    monkeypatch.setattr(morning_check, "is_nyse_trading_day", mock_is_trading)

    assert morning_check.next_morning_check_time(_t(2026, 5, 4, 18, 0)) == _t(2026, 5, 6, 7, 0)


def test_next_skips_weekend_then_holiday_monday(monkeypatch):
    """Fri evening + Mon May 11 is a holiday → roll to Tue May 12."""
    import morning_check

    real = morning_check.is_nyse_trading_day

    def mock_is_trading(d):
        if d == date(2026, 5, 11):
            return False
        return real(d)

    monkeypatch.setattr(morning_check, "is_nyse_trading_day", mock_is_trading)

    assert morning_check.next_morning_check_time(_t(2026, 5, 8, 8, 0)) == _t(2026, 5, 12, 7, 0)


# ── format_morning_check_message: silence vs alert ─────────────────

def test_format_empty_offline_returns_none():
    """All gateways online → no message (caller stays silent)."""
    from morning_check import format_morning_check_message
    assert format_morning_check_message([]) is None


def test_format_single_offline_names_account():
    from morning_check import format_morning_check_message
    msg = format_morning_check_message(["Rosa"])
    assert msg is not None
    assert "Rosa" in msg
    assert "\U0001F534" in msg            # 🔴
    assert "offline" in msg.lower()
    assert "Morning check" in msg


def test_format_multiple_offline_lists_all():
    from morning_check import format_morning_check_message
    msg = format_morning_check_message(["Rosa", "Taylor"])
    assert "Rosa" in msg
    assert "Taylor" in msg
    assert msg.count("\U0001F534") == 2   # one red dot per offline account


def test_format_includes_actionable_hint():
    """The alert should tell the user a manual Gateway restart is needed."""
    from morning_check import format_morning_check_message
    msg = format_morning_check_message(["Rosa"])
    assert "restart" in msg.lower()


# ── MORNING_SETTLE_SEC: must outlast a cold connect attempt ────────

def test_settle_exceeds_connect_timeout():
    """The followup must wait longer than gateway_monitor's connect timeout,
    or a slow-but-recovering account would be falsely reported offline."""
    from morning_check import MORNING_SETTLE_SEC
    from gateway_monitor import CONNECT_TIMEOUT_SEC
    assert MORNING_SETTLE_SEC > CONNECT_TIMEOUT_SEC


# ── constituents_staleness_warning: fresh vs stale vs unreadable ───

def test_staleness_fresh_returns_none():
    """Snapshot younger than the max age → silent."""
    from morning_check import constituents_staleness_warning
    assert constituents_staleness_warning(date(2026, 7, 3), date(2026, 7, 5)) is None


def test_staleness_exactly_max_age_returns_none():
    """Boundary: age == max_age_days is still fresh (one refresh cycle =
    7 calendar days; warn only from the first *missed* week onward)."""
    from morning_check import constituents_staleness_warning
    assert constituents_staleness_warning(
        date(2026, 6, 28), date(2026, 7, 5), max_age_days=7) is None


def test_staleness_over_max_age_warns_with_age_and_date():
    from morning_check import constituents_staleness_warning
    msg = constituents_staleness_warning(date(2026, 6, 1), date(2026, 7, 5))
    assert msg is not None
    assert "34 天" in msg                       # age spelled out
    assert "2026-06-01" in msg                  # last snapshot named
    assert "refresh_constituents" in msg        # actionable command


def test_staleness_none_snapshot_warns_unreadable():
    """Halt-preference (ADR-0003 convention): unreadable file alerts rather
    than staying silent."""
    from morning_check import constituents_staleness_warning
    msg = constituents_staleness_warning(None, date(2026, 7, 5))
    assert msg is not None
    assert "無法讀取" in msg
    assert "refresh_constituents" in msg


def test_staleness_warning_never_claims_to_autofetch():
    """ADR-0002 keeps refresh on-demand — the reminder must ask the operator
    to run it, not imply the system will."""
    from morning_check import constituents_staleness_warning
    msg = constituents_staleness_warning(date(2026, 6, 1), date(2026, 7, 5))
    assert "python data/refresh_constituents.py" in msg


# ── read_constituents_last_date: xlsx I/O companion ────────────────

def _const_workbook():
    """Minimal ADR-0002-shaped workbook: metadata sheets + digit year sheets
    (same layout SelectionEngine._load_constituents expects)."""
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.title = "Summary"
    wb.active.append(["Metric", "Value"])
    return wb


def test_read_last_date_returns_newest_snapshot_of_newest_year(tmp_path):
    from morning_check import read_constituents_last_date
    wb = _const_workbook()
    ws25 = wb.create_sheet("2025")
    ws25.append(["Date", "T1"])
    ws25.append([datetime(2025, 12, 31), "AAPL"])
    ws26 = wb.create_sheet("2026")
    ws26.append(["Date", "T1"])
    ws26.append([datetime(2026, 6, 29), "AAPL"])
    ws26.append([datetime(2026, 7, 1), "AAPL"])
    p = tmp_path / "const.xlsx"
    wb.save(p)
    assert read_constituents_last_date(p) == date(2026, 7, 1)


def test_read_last_date_ignores_metadata_sheets(tmp_path):
    """'Summary'/'Counts' are not snapshots — only digit-named year sheets
    count (mirrors the loader's `s.isdigit()` filter)."""
    from morning_check import read_constituents_last_date
    wb = _const_workbook()
    counts = wb.create_sheet("Counts")
    counts.append([datetime(2099, 1, 1), "not a snapshot"])
    ws = wb.create_sheet("2026")
    ws.append(["Date", "T1"])
    ws.append([datetime(2026, 6, 1), "AAPL"])
    p = tmp_path / "const.xlsx"
    wb.save(p)
    assert read_constituents_last_date(p) == date(2026, 6, 1)


def test_read_last_date_missing_file_returns_none(tmp_path):
    from morning_check import read_constituents_last_date
    assert read_constituents_last_date(tmp_path / "nope.xlsx") is None


def test_read_last_date_no_year_sheets_returns_none(tmp_path):
    """Metadata-only / malformed workbook → None → caller emits the
    unreadable warning instead of a bogus age."""
    from morning_check import read_constituents_last_date
    wb = _const_workbook()
    p = tmp_path / "empty.xlsx"
    wb.save(p)
    assert read_constituents_last_date(p) is None
